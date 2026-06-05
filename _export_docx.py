import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document

doc = Document('D:/allin/elde/1.docx')
with open('D:/allin/elde/docx_full.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'[{i}] {p.text}\n')
    for ti, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {ti} ===\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            line = '  Row' + str(ri) + ': ' + ' | '.join(cells)
            f.write(line + '\n')
print('Done')
